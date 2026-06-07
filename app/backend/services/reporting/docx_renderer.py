from __future__ import annotations

import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from io import BytesIO
from typing import Any

from docx import Document  # type: ignore[import-untyped]
from docx.enum.text import WD_BREAK  # type: ignore[import-untyped]
from docx.oxml import OxmlElement  # type: ignore[import-untyped]
from docx.oxml.ns import qn  # type: ignore[import-untyped]
from docx.shared import Inches, Pt  # type: ignore[import-untyped]


BLOCK_TAGS = {
    "address",
    "article",
    "aside",
    "blockquote",
    "div",
    "figure",
    "footer",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "header",
    "li",
    "main",
    "nav",
    "ol",
    "p",
    "pre",
    "section",
    "table",
    "ul",
}
SKIP_TAGS = {"script", "style", "meta", "link", "head", "title"}
INLINE_BOLD = {"b", "strong", "th"}
INLINE_ITALIC = {"i", "em", "cite"}


@dataclass(slots=True)
class TextNode:
    text: str


@dataclass(slots=True)
class ElementNode:
    tag: str
    attrs: dict[str, str] = field(default_factory=dict)
    children: list["ElementNode | TextNode"] = field(default_factory=list)


class _DocumentParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = ElementNode("document")
        self._stack: list[ElementNode] = [self.root]

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        node = ElementNode(tag.lower(), {key.lower(): value or "" for key, value in attrs})
        self._stack[-1].children.append(node)
        if tag.lower() not in {"br", "hr", "img", "input", "meta", "link"}:
            self._stack.append(node)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        for index in range(len(self._stack) - 1, 0, -1):
            if self._stack[index].tag == tag:
                del self._stack[index:]
                return

    def handle_data(self, data: str) -> None:
        if data:
            self._stack[-1].children.append(TextNode(data))


def render_docx(report_html: str) -> bytes:
    """Render report HTML into a native DOCX document.

    This intentionally avoids the old Word-compatible HTML download path.
    Returning a real Office Open XML package prevents Word from guessing the
    wrong encoding and gives tables/headings native Word structure.
    """

    document = Document()
    _configure_document(document)
    root = _parse(report_html)
    body = _first(root, "body") or root
    title = _document_title(root)
    if title and not _contains_tag(body, "h1"):
        document.add_heading(title, level=1)

    rendered_any = False
    for child in body.children:
        rendered_any = _render_block(document, child) or rendered_any

    if not rendered_any:
        document.add_paragraph(_collect_text(body) or "Report")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _parse(report_html: str) -> ElementNode:
    parser = _DocumentParser()
    parser.feed(report_html)
    parser.close()
    return parser.root


def _configure_document(document: Any) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Heading 1"].font.size = Pt(20)
    styles["Heading 2"].font.size = Pt(15)
    styles["Heading 3"].font.size = Pt(12.5)


def _render_block(document: Any, node: ElementNode | TextNode) -> bool:
    if isinstance(node, TextNode):
        text = _collapse(node.text)
        if text:
            document.add_paragraph(text)
            return True
        return False

    if node.tag in SKIP_TAGS:
        return False
    if node.tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        text = _collect_text(node)
        if not text:
            return False
        level = min(max(int(node.tag[1]), 1), 3)
        document.add_heading(text, level=level)
        return True
    if node.tag == "p":
        paragraph = document.add_paragraph()
        _render_inline(paragraph, node.children)
        return bool(paragraph.text.strip())
    if node.tag == "br":
        document.add_paragraph()
        return True
    if node.tag == "hr":
        paragraph = document.add_paragraph()
        paragraph.add_run("─" * 36)
        return True
    if node.tag == "pre":
        paragraph = document.add_paragraph()
        run = paragraph.add_run(_raw_text(node).strip())
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        return bool(paragraph.text.strip())
    if node.tag == "blockquote":
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.24)
        _render_inline(paragraph, node.children)
        return bool(paragraph.text.strip())
    if node.tag == "ul":
        return _render_list(document, node, style="List Bullet")
    if node.tag == "ol":
        return _render_list(document, node, style="List Number")
    if node.tag == "table":
        return _render_table(document, node)
    if node.tag in {"img", "svg", "canvas"}:
        paragraph = document.add_paragraph()
        paragraph.add_run("[Chart or image available in the HTML/PDF export]")
        return True
    if node.tag == "li":
        paragraph = document.add_paragraph(style="List Bullet")
        _render_inline(paragraph, node.children)
        return bool(paragraph.text.strip())

    child_blocks = [child for child in node.children if isinstance(child, ElementNode) and child.tag in BLOCK_TAGS]
    if child_blocks:
        rendered = False
        for child in node.children:
            rendered = _render_block(document, child) or rendered
        return rendered

    text = _collect_text(node)
    if text:
        document.add_paragraph(text)
        return True
    return False


def _render_list(document: Any, node: ElementNode, *, style: str) -> bool:
    rendered = False
    for child in node.children:
        if isinstance(child, ElementNode) and child.tag == "li":
            paragraph = document.add_paragraph(style=style)
            _render_inline(paragraph, child.children)
            rendered = bool(paragraph.text.strip()) or rendered
    return rendered


def _render_table(document: Any, node: ElementNode) -> bool:
    rows = []
    header_flags = []
    for row_node in _descendants(node, "tr"):
        cells = [cell for cell in row_node.children if isinstance(cell, ElementNode) and cell.tag in {"td", "th"}]
        if not cells:
            cells = [cell for cell in _descendants(row_node, "td")] + [cell for cell in _descendants(row_node, "th")]
        if cells:
            rows.append([_collect_text(cell) for cell in cells])
            header_flags.append(any(cell.tag == "th" for cell in cells))

    if not rows:
        return False

    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.autofit = True

    for row_index, row in enumerate(rows):
        word_row = table.rows[row_index]
        if row_index == 0 and header_flags[row_index]:
            _repeat_table_header(word_row)
        for column_index in range(column_count):
            cell = word_row.cells[column_index]
            cell.text = row[column_index] if column_index < len(row) else ""
            if row_index == 0 and header_flags[row_index]:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    document.add_paragraph()
    return True


def _render_inline(paragraph: Any, children: list[ElementNode | TextNode], *, bold: bool = False, italic: bool = False) -> None:
    for child in children:
        if isinstance(child, TextNode):
            text = _collapse(child.text)
            if text:
                run = paragraph.add_run(text)
                run.bold = bold
                run.italic = italic
            continue
        if child.tag in SKIP_TAGS:
            continue
        if child.tag == "br":
            paragraph.add_run().add_break(WD_BREAK.LINE)
            continue
        if child.tag in BLOCK_TAGS and child.tag not in {"span", "a", "strong", "b", "em", "i", "code"}:
            text = _collect_text(child)
            if text:
                run = paragraph.add_run(text)
                run.bold = bold or child.tag in INLINE_BOLD
                run.italic = italic or child.tag in INLINE_ITALIC
            continue
        next_bold = bold or child.tag in INLINE_BOLD
        next_italic = italic or child.tag in INLINE_ITALIC
        _render_inline(paragraph, child.children, bold=next_bold, italic=next_italic)


def _first(node: ElementNode, tag: str) -> ElementNode | None:
    for child in node.children:
        if isinstance(child, ElementNode):
            if child.tag == tag:
                return child
            found = _first(child, tag)
            if found is not None:
                return found
    return None


def _descendants(node: ElementNode, tag: str) -> list[ElementNode]:
    found: list[ElementNode] = []
    for child in node.children:
        if isinstance(child, ElementNode):
            if child.tag == tag:
                found.append(child)
            found.extend(_descendants(child, tag))
    return found


def _contains_tag(node: ElementNode, tag: str) -> bool:
    return _first(node, tag) is not None


def _document_title(root: ElementNode) -> str:
    title = _first(root, "title")
    return _collect_text(title) if title is not None else ""


def _collect_text(node: ElementNode | TextNode) -> str:
    return _collapse(_raw_text(node))


def _raw_text(node: ElementNode | TextNode) -> str:
    if isinstance(node, TextNode):
        return node.text
    return "".join(_raw_text(child) for child in node.children if not (isinstance(child, ElementNode) and child.tag in SKIP_TAGS))


def _collapse(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
